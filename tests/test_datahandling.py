import json, pytest, csv, os
from openpyxl import load_workbook
from utils.jsonhandling import jsonhandling1
from dotenv import load_dotenv


# def test_data():
#     with open("testdata/creds.json") as data:
#         finaldata=json.load(data)
#         print(finaldata["positive"]["iphone1"])
      
       
# def testdata1():
#     fdata=jsonhandling1("testdata\\creds.json")        
#     print(fdata)
    
   
def test_csvhandling():
        with open("testdata/Report (1).csv") as data:
            finaldata=csv.DictReader(data)
            output=[]
            for i in finaldata:
                output.append(i)
            print(output)    
            
           
def test_ddcsv():
    with open ("testdata/Report (1) - Report (1).csv.csv", mode="w",newline='') as data:
        finaldata=csv.DictWriter(data, fieldnames=['School Year','Course Number'])
        finaldata.writerow({'School Year': "9999", 'Course Number': "99999"})
    
              
def test_excel():
    
    workbook=load_workbook("testdata/Move (1).xlsx")
    sheet=workbook['Sheet1']                
    output=[]
    for i in sheet.iter_rows(min_row=2,values_only=True):
        output.append(i)
    print(output)    
        
                
def test_exceladd():
    
    workbook=load_workbook("testdata/Move (1).xlsx")   
    sheet=workbook['Sheet1'] 
    sheet.append(["1211","4333"])
    workbook.save("testdata/Move (1).xlsx")   
  
  
  
@pytest.mark.datahandling     
def test_CLI():
    load_dotenv(".env", override=True)
    
    #load_dotenv(os.getenv("envfile"))
    #usname=os.getenv("username","testuser")
    usname11=os.getenv("username", "pwd") 
    print(" i am CLI")
    print(usname11)

def test_json_file_read():
    data = jsonhandling1("testdata/creds.json")
    assert data["positive"]["iphone1"] == "iphone7"
    assert data["positive"]["iphone2"] == "iphone8"
